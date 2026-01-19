from __future__ import annotations

from datetime import datetime
from typing import Optional

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from src.domain.models import DossierViewModel, UiItem

def _safe_text(x: str) -> str:
    return (x or "").strip()

def export_dossier_to_word(
    vm: DossierViewModel,
    output_path: str,
    date_from: str,
    date_to: str,
    titulo: str = "Dossier Legislativo",
    cliente: Optional[str] = None,
    ) -> None:
    """
    Genera un Word con:
    - Encabezado con rango y fecha de exportación
    - Secciones por Cámara
    - Subsecciones por Apartado
    - Tabla por apartado con items
    """
    doc = Document()

    ### Estilos Base ###
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # Portada Simple
    h = doc.add_paragraph()
    run = h.add_run(titulo)
    run.bold = True
    run.font.size = Pt(20)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"Rango: {date_from} a {date_to}\n").bold = True
    if cliente:
        meta.add_run(f"Cliente: {cliente}\n")
    meta.add_run(f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}")

    doc.add_page_break()

    # Contenido
    def add_item_table(items: list[UiItem]):
        # Tabla: Fecha | Título | Promovente | Estatus | Link
        table = doc.add_table(rows=1, cols=5)
        hdr = table.rows[0].cells
        hdr[0].text = "Fecha"
        hdr[1].text = "Título"
        hdr[2].text = "Promovente"
        hdr[3].text = "Estatus"
        hdr[4].text = "Enlace"

        for it in items:
            row = table.add_row().cells
            row[0].text = _safe_text(it.fecha)
            row[1].text = _safe_text(it.titulo)
            row[2].text = _safe_text(it.promovente)
            row[3].text = _safe_text(it.estatus)
            row[4].text = _safe_text(it.link)

        doc.add_paragraph()  # espacio
    
    # Orden preferido de cámaras
    for camara in ["SENADO", "DIPUTADOS"]:
        chamber_sections = vm.get(camara) or {}
        total_cam = sum(len(v) for v in chamber_sections.values())

        # Si no hay nada en esa cámara, puedes omitirla
        if total_cam == 0:
            continue

        # Título cámara
        p = doc.add_paragraph()
        r = p.add_run(camara.title())
        r.bold = True
        r.font.size = Pt(16)

        doc.add_paragraph(f"Total de elementos: {total_cam}")
        doc.add_paragraph()

        # Orden de apartados: primero los más comunes
        preferred = [
            "COMUNICACIONES",
            "DICTAMENES",
            "DICTÁMENES",
            "INICIATIVAS",
            "PROPOSICIONES",
            "ACUERDOS PARLAMENTARIOS",
            "MINUTAS",
            "INTERVENCIONES VERBALES",
            "OTROS",
        ]
        keys = list(chamber_sections.keys())
        keys_sorted = [k for k in preferred if k in keys] + [k for k in keys if k not in preferred]

        for apartado in keys_sorted:
            items = chamber_sections.get(apartado, [])
            if not items:
                continue

            # Título apartado
            ap = doc.add_paragraph()
            rr = ap.add_run(f"{apartado.title()} ({len(items)})")
            rr.bold = True
            rr.font.size = Pt(13)

            add_item_table(items)

            # Detalle por item (resumen) - opcional, pero útil
            # Si lo quieres más compacto, comenta este bloque
            for i, it in enumerate(items, start=1):
                doc.add_paragraph(f"{i}. {it.titulo}", style="List Number")
                if it.fecha or it.estatus:
                    doc.add_paragraph(f"Fecha: {it.fecha} | Estatus: {it.estatus}")
                if it.promovente:
                    doc.add_paragraph(f"Promovente: {it.promovente}")
                if it.link:
                    doc.add_paragraph(f"Enlace: {it.link}")
                if it.resumen:
                    doc.add_paragraph(f"Resumen: {it.resumen}")
                doc.add_paragraph()

            doc.add_page_break()

    doc.save(output_path)